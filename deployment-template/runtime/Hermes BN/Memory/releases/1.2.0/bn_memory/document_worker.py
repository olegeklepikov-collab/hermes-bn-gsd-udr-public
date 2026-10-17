"""Addressed local decoders. No model, macro evaluation, arbitrary SQL, or network."""
import sys,json,math,csv,sqlite3,zipfile,hashlib,base64,os
from pathlib import Path
BLOCK=65536

def text_units(path):
    size=path.stat().st_size
    for start in range(0,max(size,1),BLOCK):
        with path.open('rb') as f:
            f.seek(start);raw=f.read(BLOCK+4)
        skip=0
        if start:
            while skip<len(raw) and raw[skip]&0xC0==0x80:skip+=1
        end=min(BLOCK,len(raw))
        while end<len(raw) and raw[end]&0xC0==0x80:end+=1
        yield {'locator':{'byte_start':start+skip,'byte_end':start+end},'text':raw[skip:end].decode('utf-8'),'extent':'original_utf8_text'}

def plain(value):
    if isinstance(value,bytes):return {'type':'binary','base64':base64.b64encode(value).decode()}
    if isinstance(value,float) and not math.isfinite(value):return {'type':'special_float','value':repr(value)}
    if value is None or isinstance(value,(bool,int,float,str)):return value
    return {'type':type(value).__name__,'representation':str(value)}

def rows_unit(rows,locator):
    return {'locator':locator,'text':json.dumps(rows,ensure_ascii=False,allow_nan=False),'extent':'exact_typed_records'}

def units(path,fmt):
    if fmt in {'text','json','xml','html','sql'}:yield from text_units(path);return
    if fmt=='rtf':
        from striprtf.striprtf import rtf_to_text
        # Preserve the complete original alongside decoded text; no field or image evaluation.
        yield from text_units(path)
        yield {'locator':{'part':'decoded_text'},'text':rtf_to_text(path.read_text(encoding='utf-8')),'extent':'rtf_decoded_text','unhandled_visuals':any(marker in path.read_text(encoding='utf-8') for marker in ('\\pict','\\object'))}
        return
    if fmt=='pdf':
        import pymupdf as fitz
        with fitz.open(path) as doc:
            if doc.needs_pass:raise ValueError('encrypted_pdf_requires_owner_access')
            for i,page in enumerate(doc):
                text=page.get_text(sort=False)
                visual=bool(page.get_images() or page.get_drawings() or not text.strip())
                row={'locator':{'page':i+1},'text':text,'extent':'extracted_page_text','visual_required':visual}
                if visual:
                    if page.rect.width*page.rect.height*2.25>20000000:raise ValueError('visual_requires_addressed_tiles')
                    raw=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).tobytes('png')
                    row['image']=base64.b64encode(raw).decode();row['image_sha256']=hashlib.sha256(raw).hexdigest()
                yield row
        return
    if fmt in {'csv','tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.reader(f,delimiter='\t' if fmt=='tsv' else ',');batch=[];start=1
            for i,row in enumerate(reader,1):
                batch.append(row)
                if len(batch)==100:yield rows_unit(batch,{'row_start':start,'row_end':i});batch=[];start=i+1
            if batch:yield rows_unit(batch,{'row_start':start,'row_end':i})
        return
    if fmt=='xlsx':
        import openpyxl
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
        values=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
        with zipfile.ZipFile(path) as package:
            extras=[n for n in package.namelist() if n.startswith(('xl/media/','xl/charts/','xl/embeddings/','xl/pivot','xl/externalLinks/','xl/comments','xl/threadedComments/','xl/connections','xl/queryTables/','customXml/')) or n.endswith('vbaProject.bin')]
        try:
            if extras:yield {'locator':{'parts':extras},'text':json.dumps(extras),'extent':'unhandled_spreadsheet_parts','unhandled_visuals':True}
            for sheet in workbook:
                cached=values[sheet.title];sheet.reset_dimensions();cached.reset_dimensions();batch=[];start=1
                for i,(row,caches) in enumerate(zip(sheet.iter_rows(),cached.iter_rows()),1):
                    batch.append([{'coordinate':__import__('openpyxl').utils.get_column_letter(j+1)+str(i),'value':plain(c.value),'cached':plain(v.value),'type':c.data_type} for j,(c,v) in enumerate(zip(row,caches))])
                    if len(batch)==100:yield rows_unit(batch,{'sheet':sheet.title,'sheet_state':sheet.sheet_state,'row_start':start,'row_end':i});batch=[];start=i+1
                if batch:yield rows_unit(batch,{'sheet':sheet.title,'sheet_state':sheet.sheet_state,'row_start':start,'row_end':i})
        finally:workbook.close();values.close()
        return
    if fmt=='xls':
        import xlrd
        book=xlrd.open_workbook(path,on_demand=True)
        try:
            for sheet in book.sheets():
                for start in range(0,sheet.nrows,100):yield rows_unit([sheet.row_values(i) for i in range(start,min(start+100,sheet.nrows))],{'sheet':sheet.name,'row_start':start+1,'row_end':min(start+100,sheet.nrows),'formula_extent':'cached_values_only'}) | {'unhandled_visuals':True}
        finally:book.release_resources()
        return
    if fmt=='docx':
        from defusedxml.ElementTree import fromstring
        with zipfile.ZipFile(path) as archive:
            names=[n for n in archive.namelist() if n.endswith(('.xml','.rels'))]
            for name in names:
                element=fromstring(archive.read(name));text=json.dumps([{'tag':x.tag,'attributes':x.attrib,'text':x.text,'tail':x.tail} for x in element.iter()],ensure_ascii=False)
                yield {'locator':{'part':name},'text':text,'extent':'all_structural_ooxml_parts','unhandled_visuals':any(n.startswith(('word/media/','word/embeddings/')) for n in archive.namelist())}
        return
    if fmt=='ods':
        from defusedxml.ElementTree import fromstring
        with zipfile.ZipFile(path) as archive:
            root=fromstring(archive.read('content.xml'))
            names=archive.namelist();visuals=[n for n in names if n.startswith(('Pictures/','Object','Thumbnails/'))]
            for name in names:
                if name.endswith('.xml') and name!='content.xml':
                    element=fromstring(archive.read(name));text=json.dumps([{'tag':x.tag,'attributes':x.attrib,'text':x.text,'tail':x.tail} for x in element.iter()],ensure_ascii=False)
                    yield {'locator':{'part':name},'text':text,'extent':'ods_structural_part','unhandled_visuals':bool(visuals)}
            if visuals:yield {'locator':{'parts':visuals},'text':json.dumps(visuals),'extent':'unhandled_ods_objects','unhandled_visuals':True}
        ns={'t':'urn:oasis:names:tc:opendocument:xmlns:table:1.0'}
        for table in root.findall('.//t:table',ns):
            for index,row in enumerate(table.findall('t:table-row',ns),1):
                values=[{'attributes':c.attrib,'text':''.join(c.itertext())} for c in row]
                yield rows_unit(values,{'table':table.attrib,'row':index})
        return
    if fmt=='sqlite':
        db=sqlite3.connect(path.as_uri()+'?mode=ro&immutable=1',uri=True);db.execute('PRAGMA query_only=ON')
        try:
            schema=db.execute("SELECT type,name,tbl_name,sql FROM sqlite_master ORDER BY type,name").fetchall()
            yield rows_unit(schema,{'kind':'complete_database_schema'})
            tables=db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name").fetchall()
            for name,ddl in tables:
                quoted='"'+name.replace('"','""')+'"';cursor=db.execute('SELECT * FROM '+quoted);cols=[x[0] for x in cursor.description]
                yield rows_unit({'columns':cols,'ddl':ddl},{'table':name,'kind':'schema'})
                start=1
                while True:
                    batch=cursor.fetchmany(100)
                    if not batch:break
                    yield rows_unit([[plain(v) for v in row] for row in batch],{'table':name,'row_start':start,'row_end':start+len(batch)-1});start+=len(batch)
        finally:db.close()
        return
    if fmt=='parquet':
        import pyarrow.parquet as pq
        file=pq.ParquetFile(path,arrow_extensions_enabled=False)
        yield rows_unit({'schema':str(file.schema),'rows':file.metadata.num_rows,'row_groups':file.metadata.num_row_groups},{'kind':'parquet_schema'})
        start=1
        for batch in file.iter_batches(batch_size=100):
            rows=[{k:plain(v) for k,v in r.items()} for r in batch.to_pylist()]
            yield rows_unit(rows,{'row_start':start,'row_end':start+len(rows)-1});start+=len(rows)
        return
    raise ValueError('format_not_supported:'+fmt)

def main():
    os.umask(0o077)
    req=json.load(sys.stdin);path=Path(req['path']);fmt=req['format'];output=Path(req['output']);output.mkdir(parents=True,exist_ok=True,mode=0o700)
    count=0;visual=0;unhandled=0
    with (output/'units.jsonl').open('w') as index:
        for i,row in enumerate(units(path,fmt)):
            text=row.pop('text');raw=text.encode();(output/(str(i)+'.txt')).write_bytes(raw)
            image=row.pop('image',None)
            if image:(output/(str(i)+'.png')).write_bytes(base64.b64decode(image));visual+=1
            if row.get('unhandled_visuals'):unhandled+=1
            row.update(index=i,characters=len(text),text_sha256=hashlib.sha256(raw).hexdigest())
            index.write(json.dumps(row,ensure_ascii=False)+'\n');count+=1
    print(json.dumps({'status':'extracted','units':count,'visual_units':visual,'unhandled_units':unhandled,'model_read':False,'scientific_acceptance':False}))

if __name__=='__main__':
    try:main()
    except Exception as e:print(json.dumps({'status':'blocked','code':type(e).__name__+':'+str(e)[:200]}));sys.exit(2)
