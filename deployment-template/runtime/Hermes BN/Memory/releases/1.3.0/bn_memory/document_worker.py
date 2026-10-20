"""Addressed local decoders. No model, macro evaluation, arbitrary SQL, or network."""
import sys,json,math,csv,sqlite3,zipfile,hashlib,base64,os,re,posixpath,subprocess,tempfile,shutil
from pathlib import Path
BLOCK=65536
FORMAT_REVISIONS={n:3 for n in ['text','json','xml','html','sql','pdf','csv','tsv','xlsx','xls','docx','ods','sqlite','parquet','rtf','pptx','odt','odp','epub','image']}

def xls_units(path,java):
    runtime=Path(__file__).resolve().parents[1]/'xls-runtime'
    pins=json.loads((runtime/'manifest.json').read_text())
    for name,digest in pins.items():
        target=runtime/name
        if target.is_symlink() or hashlib.sha256(target.read_bytes()).hexdigest()!=digest:raise ValueError('xls_runtime_integrity_failed:'+name)
    if not java or not Path(java).is_file():raise ValueError('xls_native_java_not_configured')
    # The mandatory host sandbox already encloses this worker and every child.
    # Applying a second broader sandbox is unnecessary and can be rejected by macOS.
    command=[java,'-XX:-UsePerfData','-Xmx1024m','-Djava.awt.headless=true','-Djava.io.tmpdir='+tempfile.gettempdir(),'-Dlog4j2.loggerContextFactory=org.apache.logging.log4j.simple.SimpleLoggerContextFactory','-cp',str(runtime/'classes')+os.pathsep+str(runtime/'lib/*'),'XlsReader',str(path)]
    env={k:v for k,v in os.environ.items() if k not in {'JAVA_TOOL_OPTIONS','_JAVA_OPTIONS','JDK_JAVA_OPTIONS','CLASSPATH'}}
    with tempfile.TemporaryFile(mode='w+') as errors:
        process=subprocess.Popen(command,stdout=subprocess.PIPE,stderr=errors,text=True,env=env)
        try:
            for line in process.stdout:yield json.loads(line)
            process.wait()
            if process.returncode:errors.seek(0);raise ValueError('xls_passive_reader_failed:'+errors.read()[-400:])
        finally:
            if process.stdout:process.stdout.close() # The host always terminates the complete process group.

def xml_source_text(raw):
    # This representation decodes original markup, never expands entities or evaluates it.
    if raw.startswith((b'\xff\xfe',b'\xfe\xff')):encoding='utf-16'
    elif raw.startswith(b'<\x00?\x00'):encoding='utf-16-le'
    elif raw.startswith(b'\x00<\x00?'):encoding='utf-16-be'
    else:
        match=re.match(br'<\?xml[^>]*encoding=[\'\"]([^\'\"]+)',raw.removeprefix(b'\xef\xbb\xbf'))
        encoding=match.group(1).decode('ascii') if match else 'utf-8-sig'
    return raw.decode(encoding)

def sheet_paths(package):
    from defusedxml.ElementTree import fromstring
    links={x.attrib['Id']:x.attrib['Target'] for x in fromstring(package.read('xl/_rels/workbook.xml.rels')) if x.attrib.get('TargetMode')!='External'}
    result={}
    for node in fromstring(package.read('xl/workbook.xml')).iter():
        if node.tag.rsplit('}',1)[-1]!='sheet':continue
        relation=next(v for k,v in node.attrib.items() if k.rsplit('}',1)[-1]=='id')
        target=links[relation];part=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl',target))
        if not part.startswith('xl/'):raise ValueError('worksheet_part_outside_workbook')
        result[node.attrib['name']]=part
    return result

def merged_ranges(package,part):
    from defusedxml.ElementTree import iterparse
    from openpyxl.utils.cell import range_boundaries
    ranges=[]
    with package.open(part) as stream:
        for event,node in iterparse(stream,events=('end',)):
            if node.tag.rsplit('}',1)[-1]=='mergeCell':
                ref=node.attrib['ref'];range_boundaries(ref);ranges.append(ref)
            node.clear()
    return ranges

def text_units(path,start_unit=0):
    size=path.stat().st_size
    for start in range(start_unit*BLOCK,max(size,1),BLOCK):
        with path.open('rb') as f:
            f.seek(start);raw=f.read(BLOCK+4)
        skip=0
        if start:
            while skip<len(raw) and raw[skip]&0xC0==0x80:skip+=1
        end=min(BLOCK,len(raw))
        while end<len(raw) and raw[end]&0xC0==0x80:end+=1
        yield {'locator':{'byte_start':start+skip,'byte_end':start+end},'text':raw[skip:end].decode('utf-8'),'extent':'original_utf8_text'}

def plain(value):
    if isinstance(value,bytes):return {'type':'binary','base64':base64.b64encode(value).decode()}
    if isinstance(value,float) and not math.isfinite(value):return {'type':'special_float','value':repr(value)}
    if value is None or isinstance(value,(bool,int,float,str)):return value
    return {'type':type(value).__name__,'representation':str(value)}

def rows_unit(rows,locator):
    return {'locator':locator,'text':json.dumps(rows,ensure_ascii=False,allow_nan=False),'extent':'exact_typed_records'}

def req_image_limit():return 128*1024*1024

def units(path,fmt,xls_java=None,resume=None):
    resume=resume or {};start_unit=resume.get("units",0)
    if fmt in {'text','json','xml','html','sql'}:
        for row in text_units(path,start_unit):yield row
        return
    if fmt=='rtf':
        from striprtf.striprtf import rtf_to_text
        # Preserve the complete original alongside decoded text; no field or image evaluation.
        yield from text_units(path)
        yield {'locator':{'part':'decoded_text'},'text':rtf_to_text(path.read_text(encoding='utf-8')),'extent':'rtf_decoded_text','unhandled_visuals':any(marker in path.read_text(encoding='utf-8') for marker in ('\\pict','\\object'))}
        return
    if fmt=='pdf':
        import pymupdf as fitz
        cursor=resume.get('cursor') or {'page':start_unit,'tile':0}
        with fitz.open(path) as doc:
            if doc.needs_pass:raise ValueError('encrypted_pdf_requires_owner_access')
            auxiliary={'metadata':doc.metadata,'embedded_files':[doc.embfile_info(j) for j in range(doc.embfile_count())]}
            active={k:doc.xref_get_key(doc.pdf_catalog(),k) for k in ['OpenAction','AA','Names']}
            auxiliary['catalog_references']={k:v for k,v in active.items() if v[0]!='null'}
            # Embedded documents and active actions are not followed or recursively executed.
            unresolved=bool(auxiliary['embedded_files'] or auxiliary['catalog_references'])
            for i in range(cursor['page'],len(doc)):
                page=doc[i];text=page.get_text(sort=False);needs_ocr=not text.strip();visual=bool(page.get_images() or page.get_drawings() or not text.strip())
                annotations=[{'type':a.type,'info':a.info,'rectangle':list(a.rect)} for a in (page.annots() or [])]
                widgets=[{'name':w.field_name,'value':w.field_value,'type':w.field_type_string,'rectangle':list(w.rect)} for w in (page.widgets() or [])]
                if i==0:text+='\n<PDF_METADATA_DATA>'+json.dumps(auxiliary,ensure_ascii=False,default=str)+'</PDF_METADATA_DATA>'
                if annotations or widgets:text+='\n<PDF_PAGE_DATA>'+json.dumps({'annotations':annotations,'form_fields':widgets},ensure_ascii=False,default=str)+'</PDF_PAGE_DATA>'
                clips=[page.rect]
                if visual and page.rect.width*page.rect.height*2.25>1000000:
                    step=1024/1.5;clips=[fitz.Rect(x,y,min(x+step,page.rect.x1),min(y+step,page.rect.y1)) for y in range(int(page.rect.y0),int(page.rect.y1)+1,int(step)) for x in range(int(page.rect.x0),int(page.rect.x1)+1,int(step)) if x<page.rect.x1 and y<page.rect.y1]
                for tile in range(cursor['tile'] if i==cursor['page'] else 0,len(clips)):
                    row={'locator':{'page':i+1,'tile':tile,'tiles':len(clips),'rectangle':list(clips[tile])},'text':text if tile==0 else '', 'extent':'extracted_page_text','visual_required':visual,'ocr_required':needs_ocr,'unhandled_visuals':bool(i==0 and unresolved),
                         '_next_cursor':{'page':i+1,'tile':0} if tile==len(clips)-1 else {'page':i,'tile':tile+1}}
                    if visual:
                        raw=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),clip=clips[tile],alpha=False).tobytes('png');row['image']=base64.b64encode(raw).decode();row['image_sha256']=hashlib.sha256(raw).hexdigest()
                    yield row
        return
    if fmt in {'csv','tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as f:
            csv.field_size_limit(sys.maxsize)
            reader=csv.reader(f,delimiter='\t' if fmt=='tsv' else ',');batch=[];start=1
            for i,row in enumerate(reader,1):
                batch.append(row)
                if len(batch)==100:yield rows_unit(batch,{'row_start':start,'row_end':i});batch=[];start=i+1
            if batch:yield rows_unit(batch,{'row_start':start,'row_end':i})
        return
    if fmt=='xlsx':
        import openpyxl
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
        values=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
        with zipfile.ZipFile(path) as package:
            extras=[n for n in package.namelist() if n.startswith(('xl/media/','xl/charts/','xl/embeddings/','xl/pivot','xl/externalLinks/','xl/comments','xl/threadedComments/','xl/connections','xl/queryTables/','customXml/')) or n.endswith('vbaProject.bin')]
            paths=sheet_paths(package);merges={name:merged_ranges(package,part) for name,part in paths.items()}
        try:
            if extras:yield {'locator':{'parts':extras},'text':json.dumps(extras),'extent':'unhandled_spreadsheet_parts','unhandled_visuals':True}
            for sheet in workbook:
                if merges.get(sheet.title):
                    yield rows_unit({'merged_ranges':merges[sheet.title],'semantics':'Each merged range has one anchor at its upper-left cell; covered cells share the anchor and are not independent blank observations.'},{'sheet':sheet.title,'sheet_state':sheet.sheet_state,'part':paths[sheet.title],'kind':'merged_cell_geometry'})
                cached=values[sheet.title];sheet.reset_dimensions();cached.reset_dimensions();batch=[];start=1
                for i,(row,caches) in enumerate(zip(sheet.iter_rows(),cached.iter_rows()),1):
                    batch.append([{'coordinate':__import__('openpyxl').utils.get_column_letter(j+1)+str(i),'value':plain(c.value),'cached':plain(v.value),'type':c.data_type} for j,(c,v) in enumerate(zip(row,caches))])
                    if len(batch)==100:yield rows_unit(batch,{'sheet':sheet.title,'sheet_state':sheet.sheet_state,'row_start':start,'row_end':i});batch=[];start=i+1
                if batch:yield rows_unit(batch,{'sheet':sheet.title,'sheet_state':sheet.sheet_state,'row_start':start,'row_end':i})
        finally:workbook.close();values.close()
        return
    if fmt=='xls':
        yield from xls_units(path,xls_java)
        return
    if fmt in {'docx','ods','pptx','odt','odp','epub'}:
        import codecs
        cursor=resume.get('cursor') or {'member':0,'byte':0}
        with zipfile.ZipFile(path) as archive:
            members=[n for n in archive.infolist() if not n.is_dir()]
            for member_index in range(cursor['member'],len(members)):
                member=members[member_index];name=member.filename
                if name.startswith(('/', '\\')) or '..' in name.replace('\\','/').split('/'):raise ValueError('unsafe_zip_member_path')
                if member.flag_bits&1:raise ValueError('encrypted_container_requires_owner_access')
                if member.file_size>max(1,member.compress_size)*2000:raise ValueError('container_expansion_requires_explicit_review')
                textual=name.endswith(('.xml','.rels','.txt','.html','.xhtml','.opf','.ncx','.css')) or name=='mimetype'
                if textual:
                    with archive.open(member) as source:
                        head=source.read(256);encoding='utf-16' if head.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig';decoder=codecs.getincrementaldecoder(encoding)()
                        source.seek(0);position=0
                        # Seeking a compressed stream may reread its prefix, but never materializes it in RAM.
                        skip=cursor['byte'] if member_index==cursor['member'] else 0
                        while position<member.file_size:
                            raw=source.read(BLOCK);begin=position;position+=len(raw);text=decoder.decode(raw,final=position==member.file_size)
                            if position<=skip:continue
                            yield {'locator':{'part':name,'byte_start':begin,'byte_end':position},'text':text,'extent':'passive_original_markup','_next_cursor':{'member':member_index+1,'byte':0} if position==member.file_size else {'member':member_index,'byte':position}}
                        if not member.file_size:yield {'locator':{'part':name},'text':'','extent':'empty_original_part','_next_cursor':{'member':member_index+1,'byte':0}}
                elif name.lower().endswith(('.png','.jpg','.jpeg','.tif','.tiff','.bmp','.gif','.webp')):
                    import pymupdf as fitz
                    if member.file_size>req_image_limit():raise ValueError('embedded_image_requires_addressed_intake')
                    raw=archive.read(member)
                    with fitz.open(stream=raw) as images:
                        for image_page in range(len(images)):
                            page=images[image_page];step=512;clips=[fitz.Rect(x,y,min(x+step,page.rect.x1),min(y+step,page.rect.y1)) for y in range(0,int(page.rect.y1)+1,step) for x in range(0,int(page.rect.x1)+1,step) if x<page.rect.x1 and y<page.rect.y1]
                            for tile,clip in enumerate(clips):
                                combined=image_page*1000000+tile
                                if member_index==cursor['member'] and combined<cursor.get('tile',0):continue
                                image=page.get_pixmap(clip=clip,alpha=False).tobytes('png')
                                nextc={'member':member_index+1,'byte':0} if image_page==len(images)-1 and tile==len(clips)-1 else {'member':member_index,'byte':0,'tile':combined+1}
                                yield {'locator':{'part':name,'image_page':image_page,'tile':tile,'rectangle':list(clip)},'text':'','extent':'embedded_image','visual_required':True,'image':base64.b64encode(image).decode(),'image_sha256':hashlib.sha256(image).hexdigest(),'_next_cursor':nextc}
                else:
                    yield {'locator':{'part':name},'text':json.dumps({'part':name,'bytes':member.file_size,'compressed_bytes':member.compress_size}), 'extent':'unhandled_embedded_part','unhandled_visuals':True,'_next_cursor':{'member':member_index+1,'byte':0}}
        return
    if fmt=='sqlite':
        dbpath=path;immutable=True
        with path.open('rb') as f:magic=f.read(4)
        if magic.startswith(b'PK'):
            cache=Path(tempfile.gettempdir())/'sqlite-working';cache.mkdir(exist_ok=True)
            with zipfile.ZipFile(path) as archive:
                record=json.loads(archive.read('host-fileset.json'))
                if record.get('contract')!='BNRawSQLiteFileSetV1':raise ValueError('sqlite_host_fileset_required')
                allowed={'database.sqlite','database.sqlite-wal','database.sqlite-shm'}
                if any(row['name'] not in allowed for row in record['parts']):raise ValueError('unsafe_sqlite_fileset_part')
                for row in record['parts']:
                    target=cache/row['name'];h=hashlib.sha256()
                    if target.is_symlink():raise ValueError('unsafe_sqlite_working_cache')
                    if not target.exists():
                        with archive.open(row['name']) as src,target.open('wb') as dst:
                            while chunk:=src.read(4*1024*1024):dst.write(chunk);h.update(chunk)
                        if h.hexdigest()!=row['sha256']:raise ValueError('sqlite_fileset_hash_mismatch')
            dbpath=cache/'database.sqlite';immutable=False
        db=sqlite3.connect(dbpath.as_uri()+('?mode=ro&immutable=1' if immutable else '?mode=ro'),uri=True);db.execute('PRAGMA query_only=ON');db.execute('PRAGMA trusted_schema=OFF');db.enable_load_extension(False)
        cursor=resume.get('cursor') or {'table':0,'stage':'database','key':None,'row':1}
        quote=lambda n:'"'+n.replace('"','""')+'"'
        restore=lambda v:base64.b64decode(v['base64']) if isinstance(v,dict) and v.get('type')=='binary' else v
        try:
            if cursor['stage']=='database':
                schema=db.execute('SELECT type,name,tbl_name,sql FROM sqlite_master ORDER BY type,name').fetchall()
                yield {**rows_unit(schema,{'kind':'complete_database_schema'}),'_next_cursor':{'table':0,'stage':'table','key':None,'row':1}}
            tables=db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name").fetchall()
            for ti in range(cursor['table'],len(tables)):
                name,ddl=tables[ti];info=db.execute('PRAGMA table_info('+quote(name)+')').fetchall();cols=[x[1] for x in info]
                keys=[x[1] for x in sorted(info,key=lambda v:v[5]) if x[5]] if re.search(r'\bWITHOUT\s+ROWID\b',ddl or '',re.I) else [next((n for n in ['rowid','_rowid_','oid'] if n not in cols),'')]
                if not keys or not all(keys):raise ValueError('table_requires_explicit_stable_key_export')
                stage=cursor['stage'] if ti==cursor['table'] else 'table';last=cursor.get('key') if ti==cursor['table'] else None;rownum=cursor.get('row',1) if ti==cursor['table'] else 1
                if stage in {'table','database'}:yield {**rows_unit({'columns':cols,'ddl':ddl},{'table':name,'kind':'schema'}),'_next_cursor':{'table':ti,'stage':'rows','key':last,'row':rownum}}
                while True:
                    names=','.join(quote(x) for x in keys);where=' WHERE ('+names+') > ('+','.join('?' for _ in keys)+')' if last is not None else ''
                    query='SELECT '+names+',* FROM '+quote(name)+where+' ORDER BY '+names+' LIMIT 100'
                    batch=db.execute(query,[restore(x) for x in last] if last is not None else []).fetchall()
                    if not batch:break
                    last=[plain(x) for x in batch[-1][:len(keys)]];values=[[plain(x) for x in row[len(keys):]] for row in batch]
                    yield {**rows_unit(values,{'table':name,'row_start':rownum,'row_end':rownum+len(batch)-1}),'_next_cursor':{'table':ti,'stage':'rows','key':last,'row':rownum+len(batch)}}
                    rownum+=len(batch)
        finally:db.close()
        return
    if fmt=='parquet':
        import pyarrow.parquet as pq
        file=pq.ParquetFile(path,arrow_extensions_enabled=False)
        yield rows_unit({'schema':str(file.schema),'rows':file.metadata.num_rows,'row_groups':file.metadata.num_row_groups},{'kind':'parquet_schema'})
        start=1
        for batch in file.iter_batches(batch_size=100):
            rows=[{k:plain(v) for k,v in r.items()} for r in batch.to_pylist()]
            yield rows_unit(rows,{'row_start':start,'row_end':start+len(rows)-1});start+=len(rows)
        return
    raise ValueError('format_not_supported:'+fmt)

def main():
    import time,resource,itertools,mimetypes
    mimetypes.knownfiles=[];mimetypes.init() # Do not consult user/system MIME configuration.
    os.umask(0o077)
    req=json.load(sys.stdin)
    import ctypes
    sb=ctypes.CDLL('/usr/lib/libsandbox.dylib');sb.sandbox_check.restype=ctypes.c_int
    if sb.sandbox_check(os.getpid(),None,0)!=1 or sb.sandbox_check(os.getpid(),b'network-outbound',0)!=1:raise ValueError('required_parent_isolation_not_enforced')
    path=Path(req['path']);fmt=req['format'];output=Path(req['output']);output.mkdir(parents=True,exist_ok=True,mode=0o700)
    import fcntl,threading
    child_lock=(output/'child.lock').open('a');fcntl.flock(child_lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
    def parent_watch():
        while True:
            if os.getppid()!=req['_parent_pid']:os._exit(125)
            time.sleep(.25)
    threading.Thread(target=parent_watch,daemon=True).start()
    limits=req['limits'];resource.setrlimit(resource.RLIMIT_CPU,(limits['step_seconds']+2,limits['step_seconds']+3));resource.setrlimit(resource.RLIMIT_FSIZE,(limits['max_file_bytes'],limits['max_file_bytes']));resource.setrlimit(resource.RLIMIT_NOFILE,(128,128))
    checkpoint=output/'checkpoint.json';index_path=output/'units.jsonl'
    state=json.loads(checkpoint.read_text()) if checkpoint.exists() else {'units':0,'input_units':0,'fragment':0,'index_bytes':0,'visual_units':0,'unhandled_units':0,'complete':False}
    if state['complete']:
        print(json.dumps({'status':'extracted',**state,'decoder_revision':FORMAT_REVISIONS.get(fmt,3)}));return
    started=time.monotonic();iterator=units(path,fmt,req.get('xls_java'),{'units':state['input_units'],'cursor':state.get('cursor')})
    if fmt not in {'pdf','text','json','xml','html','sql','docx','ods','pptx','odt','odp','epub','sqlite'}:iterator=itertools.islice(iterator,state['input_units'],None)
    def atomic_bytes(p,raw):
        temp=p.with_name(p.name+'.pending')
        with temp.open('wb') as f:f.write(raw);f.flush();os.fsync(f.fileno())
        os.replace(temp,p)
    def commit():atomic_bytes(checkpoint,json.dumps(state,separators=(',',':')).encode())
    if not index_path.exists():index_path.touch(mode=0o600)
    emitted=0
    with index_path.open('r+b') as index:
        index.truncate(state['index_bytes']);index.seek(state['index_bytes'])
        while emitted<limits['max_units']:
            try:row=next(iterator)
            except StopIteration:state['complete']=True;commit();break
            next_cursor=row.pop('_next_cursor',None);text=row.pop('text');image=row.pop('image',None);first_index=state['units']-state['fragment']
            if fmt=='pdf' and image and row.pop('ocr_required',not text.strip()) and Path(req['ocr_binary']).is_file():
                raster=output/'ocr-input.png';atomic_bytes(raster,base64.b64decode(image))
                try:
                    r=subprocess.run([req['ocr_binary'],str(raster)],capture_output=True,text=True,timeout=min(20,limits['step_seconds']))
                    if r.returncode:raise ValueError('ocr_unavailable')
                    observations=json.loads(r.stdout)['observations'];text+='\n'+ '\n'.join(x['text'] for x in observations);row['ocr']={'status':'extracted','engine':'AppleVision','confidence_values':[x['confidence'] for x in observations]};row['extent']='ocr_text_requires_visual_verification'
                except Exception:row['ocr']={'status':'unavailable','visual_read_required':True}
            fragments=max(1,(len(text)+BLOCK-1)//BLOCK)
            for fragment_index in range(state['fragment'],fragments):
                free=os.statvfs(output)
                if free.f_bavail*free.f_frsize<limits.get('reserve_free_bytes',536870912):raise ValueError('parser_storage_reserve_required')
                i=state['units'];piece=text[fragment_index*BLOCK:(fragment_index+1)*BLOCK];raw=piece.encode();atomic_bytes(output/(str(i)+'.txt'),raw)
                entry={**row,'locator':{**row['locator'],'fragment':fragment_index,'fragments':fragments},'index':i,'producer_attempt_id':req['_attempt_id'],'characters':len(piece),'text_sha256':hashlib.sha256(raw).hexdigest()}
                entry['visual_required']=bool(row.get('visual_required') and fragment_index==0)
                if image and fragment_index==0:atomic_bytes(output/(str(i)+'.png'),base64.b64decode(image));state['visual_units']+=1
                if row.get('unhandled_visuals'):state['unhandled_units']+=1
                index.write((json.dumps(entry,ensure_ascii=False)+'\n').encode());index.flush();os.fsync(index.fileno());state['units']+=1;state['index_bytes']=index.tell();emitted+=1
                if fragment_index==fragments-1:
                    state['input_units']+=1;state['fragment']=0
                    if next_cursor is not None:state['cursor']=next_cursor
                else:state['fragment']=fragment_index+1
                commit()
                if emitted>=limits['max_units'] or time.monotonic()-started>limits['step_seconds']*.75:break
            if emitted>=limits['max_units'] or time.monotonic()-started>limits['step_seconds']*.75:break
    print(json.dumps({'status':'extracted' if state['complete'] else 'partial',**state,'decoder_revision':FORMAT_REVISIONS.get(fmt,3),'model_read':False,'scientific_acceptance':False}))

if __name__=='__main__':
    try:main()
    except Exception as e:print(json.dumps({'status':'blocked','code':type(e).__name__+':'+str(e)[:200]}));sys.exit(2)
